# -*- coding: utf-8 -*-
"""Converts Unit 1 & Unit 2 topic HTML pages into a single shareable Word document."""
import os
from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = r"G:\Leo-Workspace\MBA-Study-Guide\Strategic-Management\units"
OUT = r"G:\Leo-Workspace\MBA-Study-Guide\Strategic-Management-Unit1-2-Notes.docx"

ACCENT = RGBColor(0x0A, 0x4F, 0x4F)
ACCENT_LIGHT = RGBColor(0x0F, 0x6E, 0x6E)
GREY = RGBColor(0x4B, 0x55, 0x63)

UNIT1_TOPICS = [
    "topic1-strategy-and-strategic-management.html",
    "topic2-stakeholders.html",
    "topic3-intent-vision-mission.html",
    "topic4-purpose-and-business-definition.html",
    "topic5-goals-and-objectives.html",
    "topic6-governance-and-csr.html",
]
UNIT2_TOPICS = [
    "topic1-business-environment.html",
    "topic2-industry-analysis.html",
    "topic3-competitor-analysis.html",
    "topic4-diamond-model-strategic-groups.html",
    "topic5-internal-appraisal.html",
    "topic6-resources-and-capabilities.html",
    "topic7-core-competence-vrio.html",
    "topic8-competitive-advantage.html",
]

def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def add_heading(doc, text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    if level == 1:
        run.font.size = Pt(20)
        run.font.color.rgb = ACCENT
    elif level == 2:
        run.font.size = Pt(15)
        run.font.color.rgb = ACCENT
    else:
        run.font.size = Pt(12.5)
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    return h

def inline_text(el):
    """Extract text from an element, marking <strong>/<em> for bold/italic runs."""
    parts = []
    for node in el.children:
        if isinstance(node, NavigableString):
            parts.append((str(node), False, False))
        elif node.name in ("strong", "b"):
            parts.append((node.get_text(), True, False))
        elif node.name in ("em", "i"):
            parts.append((node.get_text(), False, True))
        else:
            parts.append((node.get_text(), False, False))
    return parts

def add_runs(paragraph, el):
    for text, bold, italic in inline_text(el):
        if not text:
            continue
        run = paragraph.add_run(text)
        run.bold = bold
        run.italic = italic

def add_table(doc, table_el):
    rows = table_el.find_all("tr")
    if not rows:
        return
    ncols = max(len(r.find_all(["th", "td"])) for r in rows)
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r in rows:
        cells_src = r.find_all(["th", "td"])
        is_header = bool(r.find("th"))
        row = t.add_row()
        for i, c in enumerate(cells_src):
            if i >= ncols:
                break
            cell = row.cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs(p, c)
            if is_header:
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                set_cell_shading(cell, "0F6E6E")
            for run in p.runs:
                run.font.size = Pt(9.5)
    doc.add_paragraph()

def add_list(doc, list_el, ordered=False):
    for li in list_el.find_all("li", recursive=False):
        style = "List Number" if ordered else "List Bullet"
        p = doc.add_paragraph(style=style)
        # handle nested content and <details> inside <li> (practice questions)
        details = li.find("details")
        nested_lists = li.find_all(["ul", "ol"], recursive=False)
        # get direct text/inline content excluding nested lists/details
        for node in li.children:
            if getattr(node, "name", None) in ("ul", "ol", "details"):
                continue
            if isinstance(node, NavigableString):
                if node.strip():
                    p.add_run(str(node))
            elif node.name in ("strong", "b"):
                r = p.add_run(node.get_text())
                r.bold = True
            elif node.name in ("em", "i"):
                r = p.add_run(node.get_text())
                r.italic = True
            else:
                p.add_run(node.get_text())
        if details:
            summary = details.find("summary")
            answer = details.find(class_="answer-box")
            ap = doc.add_paragraph()
            ap.paragraph_format.left_indent = Inches(0.4)
            r = ap.add_run("Answer: ")
            r.bold = True
            r.font.color.rgb = ACCENT_LIGHT
            if answer:
                r2 = ap.add_run(answer.get_text())
                r2.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        for nl in nested_lists:
            add_list(doc, nl, ordered=(nl.name == "ol"))

def add_dl(doc, dl_el):
    dts = dl_el.find_all("dt")
    dds = dl_el.find_all("dd")
    for dt, dd in zip(dts, dds):
        p = doc.add_paragraph()
        r = p.add_run(dt.get_text() + ": ")
        r.bold = True
        r.font.color.rgb = ACCENT
        p.add_run(dd.get_text())

def process_content(doc, content_div):
    for el in content_div.find_all(["h2", "h3", "p", "ul", "ol", "table", "dl"], recursive=False):
        if el.name == "h2":
            txt = el.get_text().strip()
            if txt.startswith(("1.", "2.", "3.", "4.", "5.")) and any(k in txt for k in
                ["Overview", "Detailed Notes", "Key Terms", "Exam-Prep", "Practice Questions"]):
                add_heading(doc, txt, level=2)
            else:
                add_heading(doc, txt, level=2)
        elif el.name == "h3":
            add_heading(doc, el.get_text().strip(), level=3)
        elif el.name == "p":
            cls = el.get("class") or []
            if "subtitle" in cls or "breadcrumb" in cls:
                continue
            if "callout" in cls:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                add_runs(p, el)
                for run in p.runs:
                    run.italic = True
                    run.font.color.rgb = GREY
            else:
                p = doc.add_paragraph()
                add_runs(p, el)
        elif el.name == "ul":
            add_list(doc, el, ordered=False)
        elif el.name == "ol":
            cls = el.get("class") or []
            add_list(doc, el, ordered=True)
        elif el.name == "table":
            add_table(doc, el)
        elif el.name == "dl":
            add_dl(doc, el)

def add_topic(doc, unit_num, topic_idx, filepath):
    html = open(filepath, encoding="utf-8").read()
    soup = BeautifulSoup(html, "html.parser")
    content = soup.find("main", class_="content")
    h1 = content.find("h1")
    title = h1.get_text().replace(f"Unit {unit_num}", "").strip()

    doc.add_page_break()
    h = doc.add_heading(level=1)
    r1 = h.add_run(f"Unit {unit_num}.{topic_idx}  ")
    r1.font.size = Pt(13)
    r1.font.color.rgb = ACCENT_LIGHT
    r2 = h.add_run(title)
    r2.font.size = Pt(20)
    r2.font.color.rgb = ACCENT

    subtitle = content.find("p", class_="subtitle")
    if subtitle:
        p = doc.add_paragraph()
        p.add_run(subtitle.get_text()).italic = True
        for run in p.runs:
            run.font.color.rgb = GREY
            run.font.size = Pt(10)

    process_content(doc, content)

def build():
    doc = Document()

    # base font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Cover page
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(150)
    r = title_p.add_run("MBA Study Guide")
    r.font.size = Pt(32)
    r.bold = True
    r.font.color.rgb = ACCENT

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub_p.add_run("Strategic Management")
    r.font.size = Pt(20)
    r.font.color.rgb = ACCENT_LIGHT

    sub2_p = doc.add_paragraph()
    sub2_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub2_p.add_run("Units 1 & 2 \u2014 Study Notes")
    r.font.size = Pt(14)
    r.font.color.rgb = GREY

    info_p = doc.add_paragraph()
    info_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info_p.paragraph_format.space_before = Pt(40)
    r = info_p.add_run(
        "Book: Strategic Management (BA4301, Anna University MBA Sem III)\n"
        "Dr. G. Pandi Selvi & Dr. M. Hemalatha, Thakur Publication\n\n"
        "Unit 1: Strategy and Process\n"
        "Unit 2: Environmental Analysis & Competitive Advantage"
    )
    r.font.size = Pt(11)
    r.font.color.rgb = GREY

    note_p = doc.add_paragraph()
    note_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_p.paragraph_format.space_before = Pt(60)
    r = note_p.add_run(
        "This is a condensed study guide, not a copy of the textbook \u2014 each topic distills the "
        "book's content into notes, key terms, an exam-prep summary, and practice questions with "
        "model answers."
    )
    r.italic = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = GREY

    # Unit 1
    doc.add_page_break()
    h = doc.add_heading(level=1)
    r = h.add_run("UNIT 1: Strategy and Process")
    r.font.size = Pt(24)
    r.font.color.rgb = ACCENT

    for i, fname in enumerate(UNIT1_TOPICS, start=1):
        add_topic(doc, 1, i, os.path.join(ROOT, "unit1", fname))

    # Unit 2
    doc.add_page_break()
    h = doc.add_heading(level=1)
    r = h.add_run("UNIT 2: Environmental Analysis & Competitive Advantage")
    r.font.size = Pt(24)
    r.font.color.rgb = ACCENT

    for i, fname in enumerate(UNIT2_TOPICS, start=1):
        add_topic(doc, 2, i, os.path.join(ROOT, "unit2", fname))

    doc.save(OUT)
    print("Saved:", OUT)

if __name__ == "__main__":
    build()
